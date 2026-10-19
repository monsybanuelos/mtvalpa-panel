#!/usr/bin/env python3
"""Import MB-CELULARES.xlsx (form responses) into data/celumb.db."""

from __future__ import annotations

import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(
    "/cursor/stores/bc-01a14ca7-e7cb-7201-8575-dd029e780eef/docs/MB-CELULARES.xlsx"
)
DB_PATH = ROOT / "data" / "celumb.db"


def ser_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    return s[:10] if s else None


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except Exception:
        return None


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def imei_txt(v):
    if v is None or v == "" or v == 0 or v == 0.0:
        return None
    if isinstance(v, float):
        return str(int(v))
    return str(v).strip() or None


def stock_key(s: str) -> int:
    m = re.search(r"(\d+)", s)
    return int(m.group(1)) if m else 0


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print("pip install openpyxl", file=sys.stderr)
        return 1

    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"missing {xlsx}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Respuestas de formulario 1"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    phones: dict[str, dict] = {}
    sales: list[dict] = []

    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None and x != "" for x in r):
            continue
        d = {headers[i]: r[i] for i in range(len(headers)) if headers[i]}
        tipo = txt(d.get("Tipo de ingreso"))
        stock = txt(d.get("No. Stock")) or txt(d.get("No. de Stock"))
        if not stock:
            continue
        stock = stock.upper()

        if tipo == "Adquisición":
            cost = num(d.get("Costo "))
            price = num(d.get("Precio inicial sugerido"))
            if cost is None:
                continue
            if price is None:
                price = cost
            status_raw = txt(d.get("Vendido")) or ""
            status = "vendido" if status_raw.upper() == "VENDIDO" else "disponible"
            charger = txt(d.get("Incluye cargador")) or "No"
            if charger not in ("Si", "No", "Se incluira en venta"):
                charger = "Si" if charger.lower() in ("si", "sí", "yes", "1") else "No"
            phones[stock] = {
                "stock_no": stock,
                "acquired_at": ser_date(d.get("Fecha")) or "2026-01-01",
                "cost": cost,
                "suggested_price": price,
                "margin_suggested": price - cost,
                "includes_charger": charger,
                "brand": (txt(d.get("Marca")) or "—").strip(),
                "model": (txt(d.get("Modelo ")) or "—").strip(),
                "imei": imei_txt(d.get("IMEI")),
                "carrier": txt(d.get("Compañia")) or "Liberado",
                "condition_label": txt(d.get("Condicion estetica"))
                or "A Bueno (9-10)",
                "ram": d.get("RAM"),
                "storage": d.get("Almacenamiento"),
                "battery_pct": None,
                "comments": txt(d.get("Comentarios")),
                "status": status,
            }
            monto = num(d.get("Monto venta"))
            if monto and monto > 0:
                sales.append(
                    {
                        "stock_no": (txt(d.get("No. de Stock")) or stock).upper(),
                        "sold_at": ser_date(d.get("Fecha 2"))
                        or ser_date(d.get("Fecha"))
                        or phones[stock]["acquired_at"],
                        "amount": monto,
                        "customer": (txt(d.get("Cliente")) or "Cliente").strip(),
                        "channel": (txt(d.get("TIpo de venta")) or "Contado").strip()
                        or "Contado",
                        "extra_costs": num(d.get("Costos extraordinarios")) or 0,
                        "comments": txt(d.get("Comentarios 2")),
                        "pago1": num(d.get("Pago 1")),
                        "debtor_name": txt(d.get("Nombre")),
                        "debtor_phone": txt(d.get("Número")),
                        "liquidado": txt(d.get("Liquidado")) or "No",
                        "due_date": ser_date(d.get("Fecha límite de pago ")),
                        "faltante": num(d.get("Falta liquidar crédito")),
                        "pago2": txt(d.get("Pago 2")) or "No",
                        "pago3": txt(d.get("Pago 3")) or "No",
                        "credit_type": txt(d.get("Tipo de credito ")) or "Crédito",
                    }
                )
                phones[stock]["status"] = "vendido"
        elif tipo == "Venta":
            monto = num(d.get("Monto venta"))
            sale_stock = (txt(d.get("No. de Stock")) or stock).upper()
            if monto and monto > 0:
                sales.append(
                    {
                        "stock_no": sale_stock,
                        "sold_at": ser_date(d.get("Fecha 2"))
                        or ser_date(d.get("Fecha"))
                        or "2026-01-01",
                        "amount": monto,
                        "customer": (txt(d.get("Cliente")) or "Cliente").strip(),
                        "channel": (txt(d.get("TIpo de venta")) or "Contado").strip()
                        or "Contado",
                        "extra_costs": num(d.get("Costos extraordinarios")) or 0,
                        "comments": txt(d.get("Comentarios 2")),
                        "pago1": num(d.get("Pago 1")),
                        "debtor_name": txt(d.get("Nombre")),
                        "debtor_phone": txt(d.get("Número")),
                        "liquidado": txt(d.get("Liquidado")) or "No",
                        "due_date": ser_date(d.get("Fecha límite de pago ")),
                        "faltante": num(d.get("Falta liquidar crédito")),
                        "pago2": txt(d.get("Pago 2")) or "No",
                        "pago3": txt(d.get("Pago 3")) or "No",
                        "credit_type": txt(d.get("Tipo de credito ")) or "Crédito",
                    }
                )
                if sale_stock in phones:
                    phones[sale_stock]["status"] = "vendido"

    by_stock: dict[str, dict] = {}
    for s in sales:
        prev = by_stock.get(s["stock_no"])
        if (
            not prev
            or (s.get("debtor_name") and not prev.get("debtor_name"))
            or s["amount"] >= prev["amount"]
        ):
            by_stock[s["stock_no"]] = s
    sales = list(by_stock.values())

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys=OFF")
    for t in (
        "collections",
        "credit_accounts",
        "sales",
        "supplier_credits",
        "ledger",
        "phones",
    ):
        conn.execute(f"DELETE FROM {t}")
    conn.execute("PRAGMA foreign_keys=ON")

    cols = {r[1] for r in conn.execute("PRAGMA table_info(phones)")}
    if "margin_suggested" not in cols:
        conn.execute("ALTER TABLE phones ADD COLUMN margin_suggested REAL")
    if "battery_pct" not in cols:
        conn.execute("ALTER TABLE phones ADD COLUMN battery_pct REAL")
    scols = {r[1] for r in conn.execute("PRAGMA table_info(sales)")}
    if "margin" not in scols:
        conn.execute("ALTER TABLE sales ADD COLUMN margin REAL")

    ins_phone = """INSERT INTO phones (
      stock_no, acquired_at, cost, suggested_price, margin_suggested, includes_charger,
      brand, model, imei, carrier, condition_label, ram, storage, battery_pct, comments, status
    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"""

    available_after = 0
    ledger_rows: list[tuple] = []
    phone_ids: dict[str, int] = {}

    for stock in sorted(phones.keys(), key=stock_key):
        p = phones[stock]
        ram, storage = p["ram"], p["storage"]
        if ram is not None and not isinstance(ram, str):
            ram = str(int(ram)) if float(ram) == int(float(ram)) else str(ram)
        if storage is not None and not isinstance(storage, str):
            storage = (
                str(int(storage))
                if float(storage) == int(float(storage))
                else str(storage)
            )
        cur = conn.execute(
            ins_phone,
            (
                p["stock_no"],
                p["acquired_at"],
                p["cost"],
                p["suggested_price"],
                p["margin_suggested"],
                p["includes_charger"],
                p["brand"],
                p["model"],
                p["imei"],
                p["carrier"],
                p["condition_label"],
                ram,
                storage,
                p["battery_pct"],
                p["comments"],
                p["status"],
            ),
        )
        phone_ids[stock] = int(cur.lastrowid)
        if p["status"] == "disponible":
            available_after += 1
        ledger_rows.append(
            (
                p["acquired_at"] + "T12:00:00",
                "adquisicion",
                stock,
                -p["cost"],
                available_after if p["status"] == "disponible" else None,
                f"Adquisición {p['brand']} {p['model']}",
            )
        )

    for s in sorted(sales, key=lambda x: x["sold_at"]):
        stock = s["stock_no"]
        if stock not in phone_ids:
            continue
        pid = phone_ids[stock]
        cost = phones[stock]["cost"]
        margin = s["amount"] - cost - (s["extra_costs"] or 0)
        raw_tipo = str(s.get("channel") or "")
        if "credi" in raw_tipo.lower():
            channel = "Credito"
        elif raw_tipo in ("Contado", "Credito"):
            channel = raw_tipo
        else:
            channel = (
                "Credito"
                if (s.get("debtor_name") or s.get("faltante") is not None)
                else "Contado"
            )

        cur = conn.execute(
            """INSERT INTO sales (phone_id, sold_at, amount, customer, channel, extra_costs, margin, comments)
               VALUES (?,?,?,?,?,?,?,?)""",
            (
                pid,
                s["sold_at"],
                s["amount"],
                s["customer"],
                channel,
                s["extra_costs"] or 0,
                margin,
                s["comments"],
            ),
        )
        sale_id = int(cur.lastrowid)
        conn.execute("UPDATE phones SET status=? WHERE id=?", ("vendido", pid))
        avail = conn.execute(
            "SELECT COUNT(*) FROM phones WHERE status='disponible'"
        ).fetchone()[0]
        ledger_rows.append(
            (
                s["sold_at"] + "T12:00:00",
                "venta",
                stock,
                s["amount"],
                avail,
                f"Venta {channel} — {s['customer']}",
            )
        )

        if channel == "Credito":
            debtor = s.get("debtor_name") or s["customer"]
            principal = s["amount"]
            pago1 = s.get("pago1") or 0
            if s.get("faltante") is not None:
                balance = float(s["faltante"])
                paid = principal - balance
            else:
                paid = min(pago1, principal) if pago1 else 0
                balance = principal - paid
            liquidado = "Si" if balance <= 0 else (s.get("liquidado") or "No")
            if liquidado not in ("Si", "No", "Retrasado"):
                liquidado = "No"
            pago2 = s.get("pago2") or "No"
            pago3 = s.get("pago3") or "No"
            if pago2 not in ("Si", "No", "Retrasado"):
                pago2 = "No"
            if pago3 not in ("Si", "No", "Retrasado"):
                pago3 = "No"
            cur = conn.execute(
                """INSERT INTO credit_accounts (
                  sale_id, debtor_name, debtor_phone, credit_type, principal,
                  paid_total, balance, pago1, pago2_status, pago3_status, due_date, liquidado, comments
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    sale_id,
                    debtor,
                    s.get("debtor_phone"),
                    s.get("credit_type") or "Crédito",
                    principal,
                    paid,
                    max(0, balance),
                    pago1 if pago1 else None,
                    pago2,
                    pago3,
                    s.get("due_date"),
                    liquidado,
                    s.get("comments"),
                ),
            )
            credit_id = int(cur.lastrowid)
            if paid > 0:
                conn.execute(
                    """INSERT INTO collections (credit_id, amount, paid_at, note)
                       VALUES (?,?,?,?)""",
                    (credit_id, paid, s["sold_at"] + "T12:00:00", "Pago 1"),
                )
                ledger_rows.append(
                    (
                        s["sold_at"] + "T12:00:00",
                        "cobranza",
                        stock,
                        paid,
                        None,
                        f"Cobranza {debtor}",
                    )
                )

    for L in ledger_rows:
        conn.execute(
            """INSERT INTO ledger (at, kind, stock_no, amount, available_after, note)
               VALUES (?,?,?,?,?,?)""",
            L,
        )
    conn.commit()
    print(
        {
            "phones": len(phones),
            "disponible": sum(1 for p in phones.values() if p["status"] == "disponible"),
            "vendido": sum(1 for p in phones.values() if p["status"] == "vendido"),
            "sales": len(sales),
            "db": str(DB_PATH),
        }
    )
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
